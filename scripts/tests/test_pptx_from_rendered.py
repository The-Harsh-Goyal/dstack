"""Regression tests for pptx_from_rendered.py (issue #56).

Run: python3 -m unittest scripts/tests/test_pptx_from_rendered.py -v
Dependencies: playwright (with chromium), python-pptx
"""
import os, re, subprocess, sys, tempfile, unittest

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "..", "pptx_from_rendered.py")
FIXTURE = os.path.join(HERE, "fixtures", "mixed_text.html")
sys.path.insert(0, os.path.join(HERE, ".."))


def pptx_text(path):
    prs = Presentation(path)
    return " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


class MixedContentTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = os.path.join(cls.tmp.name, "mixed.pptx")
        r = subprocess.run([sys.executable, SCRIPT, FIXTURE, "-o", cls.out, "--selector", ".slide"],
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stderr
        cls.stdout = r.stdout
        cls.text = pptx_text(cls.out)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_bare_text_next_to_svg_kept(self):
        self.assertIn("Open Source", self.text)
        self.assertIn("+7.5 vs baseline", self.text)

    def test_bare_text_after_block_child_kept(self):
        self.assertIn("Text after block child", self.text)

    def test_text_transform_applied(self):
        self.assertIn("UPPERCASE LABEL", self.text)

    def test_fractional_opacity_reaches_shape_and_text(self):
        prs = Presentation(self.out)
        xml = "".join(sh._element.xml for s in prs.slides for sh in s.shapes)
        # 0.05 → alpha 5000; 0.5 → alpha 50000
        self.assertRegex(xml, r'<a:alpha val="5000"/>')
        self.assertRegex(xml, r'<a:alpha val="50000"/>')

    def test_default_run_reports_no_loss(self):
        self.assertNotIn("⚠️", self.stdout)
        self.assertIn("no text dropped", self.stdout)


class VerifyIndependenceTest(unittest.TestCase):
    """The reference source for verification must be independent of the extraction logic: text dropped during extraction must be reported by verification."""

    def test_verify_reports_text_missing_from_pptx(self):
        import pptx_from_rendered as m
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "a.pptx")
            subprocess.run([sys.executable, SCRIPT, FIXTURE, "-o", out, "--selector", ".slide"],
                           check=True, capture_output=True)
            prs = Presentation(out)
            for sh in list(prs.slides[0].shapes):          # deliberately delete a text box to simulate an extraction miss
                if sh.has_text_frame and "Open Source" in sh.text_frame.text:
                    sh._element.getparent().remove(sh._element)
            prs.save(out)
            ref = ["Open Source +7.5 vs baseline Text after block child Plain paragraph (control) "
                   "UPPERCASE LABEL half opacity text"]
            missing = m.verify_text(ref, out)
            self.assertEqual(len(missing), 1)
            self.assertTrue(any("opensource" in s for s in missing[0]))


if __name__ == "__main__":
    unittest.main()
